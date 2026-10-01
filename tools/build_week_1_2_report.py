from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(r"D:\Project 2 + 3\week 1-2 report.docx")
OUTPUT = Path(r"D:\Project 2 + 3\Splitwise-mini\artifacts\week 1-2 report.docx")

FONT = "Times New Roman"


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def add_style(doc, name, kind, size, bold=False, before=0, after=0,
              keep_with_next=False, outline_level=None):
    if name in [style.name for style in doc.styles]:
        style = doc.styles[name]
    else:
        style = doc.styles.add_style(name, kind)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = keep_with_next
    if kind == WD_STYLE_TYPE.PARAGRAPH:
        ppr = style._element.get_or_add_pPr()
        if outline_level is not None:
            outline = ppr.find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl")
                ppr.append(outline)
            outline.set(qn("w:val"), str(outline_level))
    return style


def add_body(doc, text, *, first_line=True, before=0, after=6, keep_with_next=False):
    p = doc.add_paragraph(style="Report Body")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1) if first_line else Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Report Heading 1" if level == 1 else "Report Heading 2")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Report Body")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(4)
    p.add_run("•\t")
    p.add_run(text)
    for run in p.runs:
        set_run_font(run)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9D9D9")


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for index, (label, width) in enumerate(zip(headers, widths_cm)):
        cell = header.cells[index]
        cell.width = Cm(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, size=11, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, (value, width) in enumerate(zip(values, widths_cm)):
            cell = row.cells[index]
            cell.width = Cm(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F7FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if index else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=10.5, bold=index == 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_numbered_problem(doc, number, title, text):
    p = doc.add_paragraph(style="Report Body")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    label = p.add_run(f"4.{number} {title}. ")
    set_run_font(label, bold=True)
    body = p.add_run(text)
    set_run_font(body)
    return p


def build():
    doc = Document(SOURCE)
    original_body_elements = [p._p.xml for p in doc.paragraphs]

    add_style(doc, "Report Body", WD_STYLE_TYPE.PARAGRAPH, 12, after=6)
    body_style = doc.styles["Report Body"]
    body_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body_style.paragraph_format.widow_control = True
    add_style(doc, "Report Heading 1", WD_STYLE_TYPE.PARAGRAPH, 14, bold=True,
              before=12, after=6, keep_with_next=True, outline_level=0)
    add_style(doc, "Report Heading 2", WD_STYLE_TYPE.PARAGRAPH, 12, bold=True,
              before=8, after=4, keep_with_next=True, outline_level=1)

    add_heading(doc, "1 Đặt vấn đề", 1)
    add_body(doc,
        "Trong các chuyến đi, hoạt động câu lạc bộ, nhóm bạn ở chung hoặc nhóm làm dự án, "
        "nhiều thành viên có thể lần lượt thanh toán các khoản khác nhau. Có khoản dành cho "
        "toàn nhóm, nhưng cũng có khoản chỉ liên quan đến một số người, chẳng hạn tiền taxi "
        "chỉ chia cho những người đi cùng. Khi số người và số khoản chi tăng, việc cộng thủ "
        "công rất dễ nhầm người tham gia, làm tròn sai hoặc bỏ sót khoản đã hoàn trả.")
    add_body(doc,
        "Cách xử lý phổ biến bằng tin nhắn, ghi chú hoặc bảng tính đòi hỏi một người tổng hợp "
        "lại toàn bộ dữ liệu. Mỗi lần sửa hóa đơn hoặc ghi nhận một khoản thanh toán, người đó "
        "phải tính lại số dư và danh sách chuyển tiền. Quá trình này mất thời gian, khó kiểm "
        "tra và dễ tạo tranh luận vì các thành viên không nhìn thấy rõ số tiền của mình được "
        "hình thành từ những khoản nào.")
    add_body(doc,
        "Vì vậy, đề tài hướng tới một website hỗ trợ ghi nhận chi phí nhóm, tính phần phải chịu "
        "của từng người và đề xuất cách hoàn tất công nợ. Hệ thống cần bảo toàn từng đồng, xử "
        "lý được hóa đơn chỉ chia cho một nhóm con, hỗ trợ nhiều người cùng thanh toán một "
        "khoản và cập nhật đúng sau khi có giao dịch được xác nhận. Bên cạnh độ chính xác, danh "
        "sách chuyển tiền phải đủ ngắn và dễ thực hiện trong thực tế.")

    add_heading(doc, "2 Mô tả tóm tắt đề tài", 1)
    add_body(doc,
        "Splitwise-mini là website quản lý chi phí chung cho một nhóm. Người dùng tạo nhóm, "
        "thêm thành viên và ghi nhận từng khoản chi gồm nội dung, số tiền, người đã thanh toán "
        "và danh sách người chịu phần chi. Mỗi hóa đơn có danh sách tham gia riêng; hệ thống "
        "không mặc định chia mọi khoản cho toàn bộ thành viên.")
    add_body(doc,
        "Phần chi có thể được nhập theo số tiền cụ thể, chia đều hoặc chia theo trọng số. Với "
        "các phép chia không hết, hệ thống phân bổ phần lẻ theo quy tắc xác định để tổng phần "
        "chi luôn bằng đúng tổng hóa đơn. Từ lịch sử khoản chi và các giao dịch đã xác nhận, "
        "hệ thống tính số dư ròng: số dư dương là số tiền thành viên còn được nhận, số dư âm "
        "là số tiền thành viên còn cần trả.")
    add_body(doc,
        "Sau khi có số dư, website cung cấp hai cách gợi ý hoàn trả. Chế độ Ít lần chuyển nhất "
        "tìm phương án giảm số giao dịch của cả nhóm. Chế độ Gom về một người cho phép chọn "
        "một thành viên làm đầu mối: những người còn nợ chuyển vào, sau đó người này phân phối "
        "cho những người còn được nhận. Gợi ý chỉ là kế hoạch; số dư chỉ thay đổi khi giao dịch "
        "thực tế được ghi nhận và xác nhận.")
    add_body(doc,
        "Hướng phát triển của đề tài được tham khảo từ các ứng dụng chia hóa đơn hiện có, đặc "
        "biệt là luồng tạo nhóm, nhập khoản chi, xem số dư và ghi nhận hoàn trả. Phần thuật toán "
        "được thiết kế độc lập để kiểm soát cách làm tròn, kiểm tra dữ liệu và cải thiện số lượt "
        "chuyển tiền so với cách ghép số dư lớn nhất theo từng bước.")

    add_heading(doc, "3 Mục tiêu và phạm vi đề tài", 1)
    add_heading(doc, "3.1 Mục tiêu tổng quát", 2)
    add_body(doc,
        "Mục tiêu của đề tài là xây dựng một website giúp nhóm người ghi nhận chi phí chung, "
        "xác định chính xác số dư của từng thành viên và tạo phương án hoàn trả dễ thực hiện. "
        "Kết quả tính phải có thể kiểm tra từ lịch sử dữ liệu, ổn định khi thay đổi thứ tự hiển "
        "thị và được tính lại khi hóa đơn hoặc trạng thái giao dịch thay đổi.")

    add_heading(doc, "3.2 Mục tiêu cụ thể", 2)
    objectives = [
        "Quản lý nhóm, thành viên, khoản chi và lịch sử hoàn trả bằng các mã định danh duy nhất.",
        "Cho phép một hoặc nhiều người thanh toán một khoản và chọn đúng những người chịu phần chi của khoản đó.",
        "Hỗ trợ chia theo số tiền, chia đều và chia theo trọng số; xử lý phần lẻ mà không làm mất hoặc tự thêm tiền.",
        "Tính số dư ròng từ toàn bộ khoản chi và chỉ các giao dịch đã được xác nhận.",
        "Sinh hai phương án gợi ý gồm Ít lần chuyển nhất và Gom về một người, đồng thời nêu rõ phạm vi tối ưu của từng phương án.",
        "Hỗ trợ hoàn trả từng phần, tính lại phần còn lại và lưu lịch sử để các thành viên có thể đối chiếu.",
        "Kiểm thử thuật toán bằng các trường hợp nhỏ có thể đối chiếu toàn bộ và các nhóm có tối đa 20 thành viên.",
    ]
    for item in objectives:
        add_bullet(doc, item)

    add_heading(doc, "3.3 Phạm vi thực hiện", 2)
    add_body(doc,
        "Trong giai đoạn của Project 2, đề tài tập trung vào quy trình chia chi phí trong một "
        "nhóm và phần gợi ý hoàn trả. Phạm vi được giới hạn như sau:", first_line=False)
    add_table(doc,
        ["Trong phạm vi", "Ngoài phạm vi hiện tại"],
        [
            ("Một đơn vị tiền cho mỗi nhóm; VND được lưu theo đơn vị đồng.",
             "Quy đổi tỷ giá và tổng hợp nhiều loại tiền trong cùng một phép tính."),
            ("Khoản chi có một hoặc nhiều người thanh toán và danh sách người chịu phần chi riêng.",
             "Quét hóa đơn bằng ảnh, đọc dữ liệu tự động hoặc kết nối trực tiếp với nhà cung cấp."),
            ("Chia theo số tiền, chia đều hoặc theo trọng số; bảo toàn tổng tiền sau làm tròn.",
             "Tính thuế, chiết khấu và phí dịch vụ theo quy tắc riêng của từng nhà hàng."),
            ("Hai cách gợi ý hoàn trả; ghi nhận trạng thái chờ, xác nhận hoặc từ chối.",
             "Tự động chuyển tiền qua ngân hàng, kiểm tra số dư tài khoản hoặc hoàn tiền tự động."),
            ("Tối ưu số lượt khi mọi người trong nhóm được phép bù trừ công nợ với nhau.",
             "Tối ưu đồng thời phí giao dịch, hạn mức ngân hàng hoặc bắt buộc chỉ trả đúng người đã thanh toán hộ."),
        ],
        [8.1, 8.1],
    )

    add_heading(doc, "4 Các bài toán thực tế cần giải quyết", 1)
    add_body(doc,
        "Từ bối cảnh sử dụng, hệ thống cần giải quyết các bài toán nghiệp vụ và tính toán sau:",
        first_line=False)
    add_numbered_problem(doc, 1, "Ghi nhận đúng người liên quan đến từng khoản chi",
        "Một nhóm có thể có 20 người nhưng một bữa ăn chỉ có 12 người, còn một chuyến taxi "
        "chỉ có 4 người. Dữ liệu của mỗi hóa đơn phải tách người thanh toán và người chịu phần "
        "chi; người ngoài danh sách không được phân bổ tiền.")
    add_numbered_problem(doc, 2, "Phân bổ chính xác khi chia không hết",
        "Ví dụ 100.000 đồng chia đều cho ba người không thể biểu diễn bằng ba số nguyên bằng "
        "nhau. Hệ thống phải phân bổ thành 33.334, 33.333 và 33.333 đồng theo một quy tắc ổn "
        "định, bảo đảm tổng cuối cùng vẫn là 100.000 đồng.")
    add_numbered_problem(doc, 3, "Xử lý nhiều người cùng thanh toán",
        "Một hóa đơn có thể được hai hoặc nhiều thành viên trả chung, và người trả có thể không "
        "nằm trong nhóm sử dụng món hoặc dịch vụ. Tổng tiền của danh sách người trả và tổng "
        "phần phải chịu đều phải khớp với giá trị hóa đơn.")
    add_numbered_problem(doc, 4, "Bù trừ công nợ toàn nhóm",
        "Việc theo dõi từng cặp người sẽ tạo nhiều khoản nợ nhỏ. Hệ thống cần quy đổi lịch sử "
        "chi tiêu thành một số dư ròng cho mỗi thành viên để người nợ có thể trả cho người được "
        "nhận phù hợp, dù hai người không cùng xuất hiện trong hóa đơn ban đầu.")
    add_numbered_problem(doc, 5, "Giảm số lượt chuyển tiền",
        "Cách ghép người còn nợ nhiều nhất với người còn được nhận nhiều nhất cho kết quả nhanh "
        "nhưng không luôn cho số lượt nhỏ nhất. Thuật toán cần tìm nghiệm chính xác cho nhóm "
        "trong ngưỡng xử lý, đồng thời trả phương án hợp lệ và không khẳng định quá mức đối với "
        "trường hợp lớn chưa chứng minh được tối ưu.")
    add_numbered_problem(doc, 6, "Gom tiền qua một thành viên",
        "Trong nhiều nhóm, mọi người muốn chuyển qua một người quen để dễ đối chiếu. Thuật toán "
        "phải xử lý được khi người gom đang nợ, đang được nhận hoặc có số dư bằng 0; thu tiền "
        "trước, phân phối sau và không tạo giao dịch với chính người gom.")
    add_numbered_problem(doc, 7, "Cập nhật sau khi hoàn trả từng phần",
        "Một người có thể trả một phần, trả thừa hoặc để giao dịch ở trạng thái chờ xác nhận. "
        "Chỉ giao dịch đã xác nhận được đưa vào số dư; khi dữ liệu thay đổi, phương án cũ phải "
        "được tính lại để tránh hướng dẫn trả trùng.")
    add_numbered_problem(doc, 8, "Đảm bảo khả năng kiểm tra và chống ghi lặp",
        "Mỗi khoản chi và mỗi giao dịch cần có mã duy nhất, lịch sử rõ ràng và quyền xác nhận "
        "phù hợp. Hệ thống phải tính lại từ một bản sổ nhất quán, không cộng cùng một giao dịch "
        "hai lần và giữ được căn cứ giải thích số dư cho từng thành viên.")

    # Confirm the original cover paragraphs were not rewritten or reformatted.
    for before, after in zip(original_body_elements, doc.paragraphs[:len(original_body_elements)]):
        if before != after._p.xml:
            raise RuntimeError("The cover page paragraphs changed")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
