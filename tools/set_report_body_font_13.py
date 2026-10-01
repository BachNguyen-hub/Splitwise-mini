from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.shared import Cm


SOURCE = Path(r"D:\Project 2 + 3\week 1-2 report.docx")
OUTPUT = Path(r"D:\Project 2 + 3\Splitwise-mini\artifacts\week 1-2 report.docx")
COVER_PARAGRAPHS = 21


def set_run_size(run, points=13):
    run.font.size = Pt(points)
    rpr = run._element.get_or_add_rPr()
    for tag in ("sz", "szCs"):
        node = rpr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            rpr.append(node)
        node.set(qn("w:val"), str(points * 2))


def main():
    doc = Document(SOURCE)
    if len(doc.paragraphs) <= COVER_PARAGRAPHS:
        raise RuntimeError("Không tìm thấy phần nội dung sau trang bìa")
    cover_text = [p.text for p in doc.paragraphs[:COVER_PARAGRAPHS]]

    for style_name in ("Report Body", "Report Heading 1", "Report Heading 2"):
        if style_name in [style.name for style in doc.styles]:
            doc.styles[style_name].font.size = Pt(13)

    for paragraph in doc.paragraphs[COVER_PARAGRAPHS:]:
        for run in paragraph.runs:
            set_run_size(run)
        if paragraph.text.startswith("4.8 "):
            # Keep the final numbered item on a clean page inside the body
            # section. Explicit zero indents are relative to the 3 cm margin.
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.right_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_size(run)

    if cover_text != [p.text for p in doc.paragraphs[:COVER_PARAGRAPHS]]:
        raise RuntimeError("Nội dung trang bìa đã thay đổi")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
