import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"D:\Project 2 + 3\week 1-2 report.docx")
OUTPUT = Path(r"D:\Project 2 + 3\Splitwise-mini\artifacts\week 1-2 report.docx")
COVER_PARAGRAPHS = 21


def set_section_type(sect_pr, value="nextPage"):
    node = sect_pr.find(qn("w:type"))
    if node is None:
        node = OxmlElement("w:type")
        sect_pr.insert(0, node)
    node.set(qn("w:val"), value)


def add_body_section(doc):
    if len(doc.sections) > 1:
        return
    body_sect_pr = doc._element.body.sectPr
    cover_sect_pr = deepcopy(body_sect_pr)
    set_section_type(cover_sect_pr)
    cover_p_pr = doc.paragraphs[COVER_PARAGRAPHS - 1]._p.get_or_add_pPr()
    existing = cover_p_pr.find(qn("w:sectPr"))
    if existing is not None:
        cover_p_pr.remove(existing)
    cover_p_pr.append(cover_sect_pr)


def replace_paragraph_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def format_paragraph(paragraph):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(6)


def main():
    doc = Document(SOURCE)
    if len(doc.paragraphs) <= COVER_PARAGRAPHS:
        raise RuntimeError("Không tìm thấy phần nội dung sau trang bìa")

    cover_text = [p.text for p in doc.paragraphs[:COVER_PARAGRAPHS]]
    add_body_section(doc)

    # The final section contains the report body; the first section preserves
    # the original cover layout.
    body_section = doc.sections[-1]
    body_section.left_margin = Cm(3)
    body_section.top_margin = Cm(2)
    body_section.right_margin = Cm(2)
    body_section.bottom_margin = Cm(2)
    body_section.start_type = WD_SECTION.NEW_PAGE

    for style_name in ("Report Body", "Report Heading 1", "Report Heading 2"):
        if style_name in [style.name for style in doc.styles]:
            fmt = doc.styles[style_name].paragraph_format
            fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            fmt.line_spacing = 1.5
            fmt.space_after = Pt(6)

    for paragraph in doc.paragraphs[COVER_PARAGRAPHS:]:
        format_paragraph(paragraph)
        if paragraph.style.name.startswith("Report Heading"):
            paragraph.paragraph_format.keep_with_next = True

    # Follow the user's sample "1. Đặt vấn đề" for top-level headings only.
    for paragraph in doc.paragraphs[COVER_PARAGRAPHS:]:
        if paragraph.style.name != "Report Heading 1":
            continue
        updated = re.sub(r"^(\d+)(?:\.)?\s+", r"\1. ", paragraph.text, count=1)
        replace_paragraph_text(paragraph, updated)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph)

    if cover_text != [p.text for p in doc.paragraphs[:COVER_PARAGRAPHS]]:
        raise RuntimeError("Nội dung trang bìa đã thay đổi")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
